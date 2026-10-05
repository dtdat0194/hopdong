# -*- coding: utf-8 -*-
"""Đọc Onboarding_List.xlsx thành dữ liệu sẵn sàng để render hợp đồng."""
import io
from datetime import timedelta

import openpyxl

from .util import doc_so_tien, ngay, s, tien

SHEET_CHINH = "Onboarding"
SHEET_CONG_TY = "CongTy"
SHEET_LUONG = "Luong"

NHAN = {
    "ho_ten": "Họ và tên", "ngay_sinh": "Ngày sinh", "noi_cu_tru": "Nơi cư trú",
    "so_cccd": "Số CCCD", "ngay_cap_cccd": "Ngày cấp CCCD", "noi_cap_cccd": "Nơi cấp CCCD",
    "chuc_danh": "Chức danh", "bo_phan": "Bộ phận", "dia_diem_lam_viec": "Địa điểm làm việc",
    "linh_vuc_cong_viec": "Lĩnh vực chuyên môn", "ngay_bat_dau": "Ngày bắt đầu",
    "ngay_ket_thuc": "Ngày kết thúc", "muc_luong": "Mức lương / Thù lao",
    "so_tai_khoan": "Số tài khoản", "ten_ngan_hang": "Tên ngân hàng",
    "chu_tai_khoan": "Chủ tài khoản", "so_hop_dong": "Số HĐ", "ngay_ky": "Ngày ký HĐ",
}

BAT_BUOC = {
    "HDLD": ["ho_ten", "ngay_sinh", "noi_cu_tru", "so_cccd", "ngay_cap_cccd", "noi_cap_cccd",
             "chuc_danh", "bo_phan", "dia_diem_lam_viec", "ngay_bat_dau", "ngay_ket_thuc",
             "muc_luong", "so_tai_khoan", "ten_ngan_hang", "chu_tai_khoan", "ngay_ky"],
    "HDKS": ["ho_ten", "ngay_sinh", "noi_cu_tru", "so_cccd", "ngay_cap_cccd", "noi_cap_cccd",
             "linh_vuc_cong_viec", "ngay_bat_dau", "ngay_ket_thuc", "muc_luong",
             "so_tai_khoan", "ten_ngan_hang", "chu_tai_khoan", "ngay_ky"],
}

MISSING = "…………………………"

# Các ô phái sinh từ ô bắt buộc (ngày ký tách thành ngày/tháng/năm, thời hạn tính từ hai mốc).
# Thiếu thì cũng phải để chỗ trống, nếu không số hợp đồng sẽ thành "015./HDKS-QCD".
PHAI_SINH = ["thang_ky", "nam_ky", "nam_hop_dong", "so_thang"]
CHO_TRONG = {"ngay_ky": "…..", "thang_ky": "…..", "nam_ky": "…….",
             "nam_hop_dong": "…….", "so_thang": "…", "so_hop_dong": "……"}


class LoiFileExcel(Exception):
    pass


class _Sheet:
    """Một sheet đọc qua hai lần mở file: ưu tiên giá trị Excel đã tính sẵn.

    Cột như 'Ngày kết thúc' thường là công thức (=EDATE(...)). Mở với data_only=True
    cho ra kết quả Excel đã lưu; ô nào chưa có kết quả thì lấy giá trị thô, và bỏ qua
    nếu đó chỉ là chuỗi công thức chưa được Excel tính lần nào.
    """

    def __init__(self, wb_val, wb_fml, ten):
        self.v, self.f = wb_val[ten], wb_fml[ten]
        self.max_row = self.f.max_row

    def __call__(self, r, c):
        val = self.v.cell(row=r, column=c).value
        if val is not None:
            return val
        raw = self.f.cell(row=r, column=c).value
        return None if isinstance(raw, str) and raw.startswith("=") else raw

    def hang(self, r):
        return self.f[r]


def _cong_ty(ws):
    cfg = {}
    for r in range(2, ws.max_row + 1):
        k = s(ws(r, 1))
        if k:
            cfg[k] = s(ws(r, 2))
    return cfg


def _luong(ws):
    """{họ tên: (lương tháng VNĐ, số giờ/tuần)} - tính lại bằng Python phòng khi file chưa có cached value."""
    try:
        ty_gia = float(ws(1, 2) or 0)
    except (TypeError, ValueError):
        ty_gia = 0
    bang = {}
    for r in range(3, ws.max_row + 1):
        ten, gio = s(ws(r, 2)), ws(r, 6)
        if not ten or not gio:
            continue
        vnd, usd = ws(r, 4), ws(r, 5)
        try:
            if vnd:
                bang[ten] = (float(vnd) * float(gio) * 4, float(gio))
            elif usd and ty_gia:
                bang[ten] = (float(usd) * float(gio) * 4 * ty_gia, float(gio))
        except (TypeError, ValueError):
            continue
    return bang


def _so_cot(row, head):
    """Giá trị ô, bỏ qua công thức chưa được Excel tính."""
    v = row.get(head)
    if isinstance(v, str) and v.startswith("="):
        return None
    return v


def build_context(row, cfg, luong):
    loai = s(row.get("Loại hợp đồng")).upper()
    ho_ten = s(row.get("Họ và tên"))
    bd, kt, ky = ngay(row.get("Ngày bắt đầu")), ngay(row.get("Ngày kết thúc")), ngay(row.get("Ngày ký HĐ"))

    # Thời hạn lấy thẳng từ cột trong Excel; thiếu thì suy ra từ hai mốc ngày.
    # Ngày kết thúc là ngày cuối cùng của kỳ hạn (= EDATE(bắt đầu, số tháng) - 1 ngày),
    # nên phải cộng lại 1 ngày trước khi đếm số tháng.
    so_thang = ""
    try:
        thang = float(_so_cot(row, "Thời hạn (tháng)") or 0)
        so_thang = f"{int(thang):02d}" if thang > 0 else ""
    except (TypeError, ValueError):
        pass
    if not so_thang and bd and kt:
        sau = kt + timedelta(days=1)
        so_thang = f"{max((sau.year - bd.year) * 12 + (sau.month - bd.month), 0):02d}"

    muc_luong, gio_tuan = luong.get(ho_ten, (None, None))
    if _so_cot(row, "Mức lương / Thù lao (VNĐ/tháng)"):
        muc_luong = float(_so_cot(row, "Mức lương / Thù lao (VNĐ/tháng)"))
    if _so_cot(row, "Số giờ/tuần"):
        gio_tuan = float(_so_cot(row, "Số giờ/tuần"))

    cccd = s(row.get("Số CCCD"))
    if cccd.isdigit() and len(cccd) < 12:
        cccd = cccd.zfill(12)

    ngay_tra = s(row.get("Ngày trả lương hằng tháng"))
    return loai, {
        "cong_ty": cfg.get("cong_ty", ""),
        "dia_chi_cong_ty": cfg.get("dia_chi_cong_ty", ""),
        "dien_thoai_cong_ty": cfg.get("dien_thoai_cong_ty", ""),
        "mst_cong_ty": cfg.get("mst_cong_ty", ""),
        "ngay_cap_dkkd": cfg.get("ngay_cap_dkkd", ""),
        "noi_cap_dkkd": cfg.get("noi_cap_dkkd", ""),
        "dai_dien": cfg.get("dai_dien", ""),
        "xung_ho_dai_dien": cfg.get("xung_ho_dai_dien", "Ông"),
        "chuc_vu_dai_dien": cfg.get("chuc_vu_dai_dien", ""),
        "ky_hieu_hd": cfg.get("ky_hieu_hdld" if loai == "HDLD" else "ky_hieu_hdks", ""),

        "ho_ten": ho_ten.upper(),
        "ngay_sinh": s(row.get("Ngày sinh")),
        "noi_cu_tru": s(row.get("Nơi cư trú")),
        "so_cccd": cccd,
        "ngay_cap_cccd": s(row.get("Ngày cấp CCCD")),
        "noi_cap_cccd": s(row.get("Nơi cấp CCCD")),

        "chuc_danh": s(row.get("Chức danh (VN)")),
        "bo_phan": s(row.get("Bộ phận")),
        "dia_diem_lam_viec": s(row.get("Địa điểm làm việc")),
        "linh_vuc_cong_viec": s(row.get("Lĩnh vực chuyên môn (HĐ dịch vụ)")),

        "so_hop_dong": s(row.get("Số HĐ")),
        "nam_hop_dong": ky.strftime("%Y") if ky else "",
        "ngay_ky": ky.strftime("%d") if ky else "",
        "thang_ky": ky.strftime("%m") if ky else "",
        "nam_ky": ky.strftime("%Y") if ky else "",
        "ngay_bat_dau": s(bd),
        "ngay_ket_thuc": s(kt),
        "so_thang": so_thang,

        "muc_luong": tien(muc_luong) if muc_luong else "",
        "muc_luong_bang_chu": s(row.get("Số tiền bằng chữ")) or (doc_so_tien(muc_luong) if muc_luong else ""),
        "gio_tuan": f"{gio_tuan:g}" if gio_tuan else "",
        "gio_ngay": f"{gio_tuan / 5:g}" if gio_tuan else "",
        "so_ngay_tuan": "05",
        "ngay_tra_luong": ngay_tra.zfill(2) if ngay_tra else "",
        "ngay_phep_nam": s(row.get("Ngày phép năm")),
        "ty_le_khau_tru": s(row.get("Tỷ lệ khấu trừ TNCN (%)")),

        "so_tai_khoan": s(row.get("Số tài khoản")),
        "ten_ngan_hang": s(row.get("Tên ngân hàng")),
        "chi_nhanh": s(row.get("Chi nhánh")),
        "chu_tai_khoan": s(row.get("Chủ tài khoản")),
    }


def cap_so_hop_dong(people):
    """Dòng nào bỏ trống 'Số HĐ' thì cấp số kế tiếp theo từng loại hợp đồng."""
    ke_tiep = {}
    for loai in BAT_BUOC:
        da_dung = [int(p["ctx"]["so_hop_dong"]) for p in people
                   if p["type"] == loai and p["ctx"]["so_hop_dong"].isdigit()]
        ke_tiep[loai] = max(da_dung) + 1 if da_dung else 1
    for p in people:
        if p["type"] and not p["ctx"]["so_hop_dong"]:
            p["ctx"]["so_hop_dong"] = f"{ke_tiep[p['type']]:03d}"
            p["contract_no"] = p["ctx"]["so_hop_dong"]
            p["auto_no"] = True
            ke_tiep[p["type"]] += 1
    return people


def thieu_thong_tin(loai, ctx):
    if loai not in BAT_BUOC:
        return []
    return [NHAN[k] for k in BAT_BUOC[loai] if not ctx.get(k)]


def dien_cho_trong(ctx, loai):
    """Thay các ô còn thiếu bằng dấu chấm lửng để bản in vẫn dùng được."""
    out = dict(ctx)
    for k in BAT_BUOC.get(loai, []) + PHAI_SINH + ["so_hop_dong"]:
        if not out.get(k):
            out[k] = CHO_TRONG.get(k, MISSING)
    return out


def doc_file(data):
    """data: đường dẫn hoặc bytes. Trả về dict gồm công ty + danh sách nhân sự."""
    blob = bytes(data) if isinstance(data, (bytes, bytearray)) else open(data, "rb").read()
    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(blob), data_only=False)
        wb_v = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    except Exception as e:
        raise LoiFileExcel(f"Không đọc được file Excel: {e}")

    if SHEET_CHINH not in wb_f.sheetnames:
        raise LoiFileExcel(
            f"File thiếu sheet '{SHEET_CHINH}'. Hãy tải lên đúng file Onboarding_List.xlsx "
            f"đã được cấu trúc lại (các sheet tìm thấy: {', '.join(wb_f.sheetnames)})."
        )
    if SHEET_CONG_TY not in wb_f.sheetnames:
        raise LoiFileExcel(f"File thiếu sheet '{SHEET_CONG_TY}'.")

    cfg = _cong_ty(_Sheet(wb_v, wb_f, SHEET_CONG_TY))
    luong = _luong(_Sheet(wb_v, wb_f, SHEET_LUONG)) if SHEET_LUONG in wb_f.sheetnames else {}
    ws = _Sheet(wb_v, wb_f, SHEET_CHINH)
    heads = {s(c.value): i for i, c in enumerate(ws.hang(2), 1) if c.value}
    if "Họ và tên" not in heads:
        raise LoiFileExcel("Sheet 'Onboarding' không có cột 'Họ và tên' ở hàng 2.")

    people = []
    for r in range(3, ws.max_row + 1):
        raw = {h: ws(r, i) for h, i in heads.items()}
        ten = s(raw.get("Họ và tên"))
        if not ten:
            continue
        loai, ctx = build_context(raw, cfg, luong)
        thieu = thieu_thong_tin(loai, ctx)
        people.append({
            "row": r,
            "stt": len(people) + 1,
            "name": ten,
            "type": loai if loai in BAT_BUOC else "",
            "type_label": {"HDLD": "HĐ lao động", "HDKS": "HĐ dịch vụ kỹ sư"}.get(loai, "Chưa chọn"),
            "role": s(raw.get("Chức danh (VN)")) or s(raw.get("Role (EN)")),
            "department": s(raw.get("Bộ phận")),
            "employment": s(raw.get("Hình thức")),
            "contract_no": ctx["so_hop_dong"],
            "sign_date": s(raw.get("Ngày ký HĐ")),
            "start": ctx["ngay_bat_dau"],
            "end": ctx["ngay_ket_thuc"],
            "months": ctx["so_thang"],
            "salary": ctx["muc_luong"],
            "salary_text": ctx["muc_luong_bang_chu"],
            "bank": " - ".join(x for x in [ctx["so_tai_khoan"], ctx["ten_ngan_hang"]] if x),
            "marked": s(raw.get("Tạo HĐ (x)")).lower() == "x",
            "missing": thieu,
            "ready": loai in BAT_BUOC and not thieu,
            "auto_no": False,
            "ctx": ctx,
        })

    cap_so_hop_dong(people)
    for p in people:
        p["detail"] = _chi_tiet(p["ctx"])
    return {"company": cfg, "people": people}


def _chi_tiet(ctx):
    so_hd = (f'{ctx["so_hop_dong"]}.{ctx["nam_hop_dong"]}/{ctx["ky_hieu_hd"]}'
             if ctx["so_hop_dong"] and ctx["nam_hop_dong"] else "")
    return [
        ("Họ và tên", ctx["ho_ten"]), ("Ngày sinh", ctx["ngay_sinh"]),
        ("Số CCCD", ctx["so_cccd"]), ("Ngày cấp CCCD", ctx["ngay_cap_cccd"]),
        ("Nơi cấp CCCD", ctx["noi_cap_cccd"]), ("Nơi cư trú", ctx["noi_cu_tru"]),
        ("Chức danh", ctx["chuc_danh"]), ("Bộ phận", ctx["bo_phan"]),
        ("Địa điểm làm việc", ctx["dia_diem_lam_viec"]),
        ("Lĩnh vực chuyên môn", ctx["linh_vuc_cong_viec"]),
        ("Số hợp đồng", so_hd),
        ("Ngày ký", "/".join(x for x in [ctx["ngay_ky"], ctx["thang_ky"], ctx["nam_ky"]] if x)),
        ("Thời hạn", f'{ctx["ngay_bat_dau"]} - {ctx["ngay_ket_thuc"]} ({ctx["so_thang"]} tháng)'
         if ctx["ngay_bat_dau"] else ""),
        ("Mức lương / Thù lao", f'{ctx["muc_luong"]} VNĐ ({ctx["muc_luong_bang_chu"]})'
         if ctx["muc_luong"] else ""),
        ("Số giờ/tuần", ctx["gio_tuan"]),
        ("Số tài khoản", ctx["so_tai_khoan"]), ("Ngân hàng", ctx["ten_ngan_hang"]),
        ("Chi nhánh", ctx["chi_nhanh"]), ("Chủ tài khoản", ctx["chu_tai_khoan"]),
    ]
