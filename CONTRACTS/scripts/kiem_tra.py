#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm tra trước khi đẩy mã lên: dữ liệu nhân sự, bản web, và việc dựng hợp đồng.

Chạy tay:  python3 CONTRACTS/scripts/kiem_tra.py
GitHub Actions chạy file này mỗi lần push, nên lỗi bị chặn lại trước khi lên trang thật.
"""
import io
import re
import subprocess
import sys
from pathlib import Path

GOC = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(GOC / "CONTRACTS"))

OK, LOI = "  ✓", "  ✗"


def git(*args):
    return subprocess.run(["git", "-C", str(GOC), *args],
                          capture_output=True, text=True).stdout.splitlines()


# ───────────────────────── 1. dữ liệu nhân sự ─────────────────────────

# Đây là kho mã công khai. Không file nào trong đây được phép chứa thông tin
# của nhân viên: không file Excel, không hợp đồng đã xuất, không script có
# sẵn họ tên và CCCD.
CAM_DUONG_DAN = [
    (re.compile(r"Onboarding_List", re.I), "danh sách onboarding"),
    (re.compile(r"(^|/)output/"), "thư mục hợp đồng đã xuất"),
    (re.compile(r"build_excel\.py$"), "script chứa dữ liệu nhân viên"),
    (re.compile(r"\.xlsx?$", re.I), "file Excel"),
]

# Dấu hiệu dữ liệu thật lọt vào mã nguồn hoặc template.
CAM_NOI_DUNG = [
    (re.compile(r"(?<!\d)\d{12}(?!\d)"), "số giống CCCD (12 chữ số)"),
    (re.compile(r"[\w.+-]+@[\w-]+\.[a-z]{2,}", re.I), "địa chỉ email"),
    (re.compile(r"(?<!\d)0\d{9}(?!\d)"), "số giống số điện thoại"),
]
# Chuỗi vô hại khớp nhầm với các mẫu trên.
BO_QUA = re.compile(r"example\.com|@v\d|dtdat0194@|\.github\.io")


def doc_chu(duong_dan):
    """Lấy phần chữ của file để soi, kể cả file .docx."""
    p = GOC / duong_dan
    if duong_dan.endswith(".docx"):
        from docx import Document
        d = Document(p)
        return "\n".join([*(x.text for x in d.paragraphs),
                          *(c.text for t in d.tables for r in t.rows for c in r.cells)])
    if duong_dan.endswith((".py", ".js", ".html", ".css", ".md", ".txt", ".yml", ".gitignore")):
        return p.read_text(encoding="utf-8", errors="ignore")
    return ""


def kiem_tra_du_lieu_ca_nhan():
    print("1. Dữ liệu nhân sự không lọt vào kho mã")
    sach = True
    theo_doi = git("ls-files")

    for duong_dan in theo_doi:
        for mau, ten in CAM_DUONG_DAN:
            if mau.search(duong_dan):
                print(f"{LOI} {duong_dan} — {ten}, không được đẩy lên")
                sach = False

    for duong_dan in theo_doi:
        chu = doc_chu(duong_dan)
        for mau, ten in CAM_NOI_DUNG:
            dinh = [m for m in mau.findall(chu) if not BO_QUA.search(m)]
            if dinh:
                print(f"{LOI} {duong_dan} — có {ten}: {dinh[:3]}")
                sach = False

    if sach:
        print(f"{OK} {len(theo_doi)} file đang theo dõi, không file nào chứa dữ liệu nhân sự")
    return sach


# ───────────────────── 2. bản web nạp đủ file ─────────────────────

def danh_sach_js(bien, js):
    """Đọc một mảng chuỗi khai báo trong backend-pyodide.js."""
    khoi = re.search(rf"const {bien}\s*=\s*\[(.*?)\]", js, re.S)
    return re.findall(r'"([^"]+)"', khoi.group(1)) if khoi else []


def kiem_tra_ban_web():
    print("2. Bản chạy trong trình duyệt nạp đủ file")
    js = (GOC / "web/backend-pyodide.js").read_text(encoding="utf-8")
    du = True

    khai_bao = {
        "CONTRACTS/app/{}.py": danh_sach_js("MODULE", js),
        "CONTRACTS/templates/{}.docx": danh_sach_js("TEMPLATE", js),
        "CONTRACTS/fonts/{}.ttf": danh_sach_js("FONT", js),
    }
    for khuon, ten_list in khai_bao.items():
        if not ten_list:
            print(f"{LOI} không đọc được danh sách cho {khuon}")
            du = False
        for ten in ten_list:
            if not (GOC / khuon.format(ten)).exists():
                print(f"{LOI} thiếu {khuon.format(ten)} (backend-pyodide.js có khai nhưng file không tồn tại)")
                du = False

    # Lỗi dễ mắc nhất: thêm module Python mới mà quên khai vào danh sách,
    # bản chạy trên máy vẫn chạy còn bản web thì chết ngay lúc khởi động.
    tren_dia = {p.stem for p in (GOC / "CONTRACTS/app").glob("*.py")} - {"server"}
    thieu = tren_dia - set(khai_bao["CONTRACTS/app/{}.py"])
    if thieu:
        print(f"{LOI} chưa khai vào MODULE của backend-pyodide.js: {sorted(thieu)}")
        du = False

    if not (GOC / ".nojekyll").exists():
        print(f"{LOI} thiếu .nojekyll — GitHub Pages sẽ bỏ qua app/__init__.py")
        du = False

    if du:
        print(f"{OK} đủ {sum(len(v) for v in khai_bao.values())} file bản web cần tải")
    return du


# ──────────────── 3. hai trang HTML có đủ phần tử app.js cần ────────────────

def kiem_tra_giao_dien():
    print("3. Hai bản giao diện có đủ phần tử app.js dùng")
    app_js = (GOC / "CONTRACTS/app/static/app.js").read_text(encoding="utf-8")
    can = set(re.findall(r'\$\("([^"]+)"\)', app_js))
    du = True
    for trang in ("index.html", "CONTRACTS/app/static/index.html"):
        co = set(re.findall(r'id="([^"]+)"', (GOC / trang).read_text(encoding="utf-8")))
        if (thieu := can - co):
            print(f"{LOI} {trang} thiếu phần tử: {sorted(thieu)}")
            du = False
    if du:
        print(f"{OK} cả hai trang đều có đủ {len(can)} phần tử")
    return du


# ──────────────── 4. dựng thử hợp đồng từ template thật ────────────────

def kiem_tra_dung_hop_dong():
    print("4. Dựng thử hợp đồng từ template")
    from docxtpl import DocxTemplate

    from app import contract_spec as spec
    from app.excel import build_context, dien_cho_trong
    from app.render_docx import TEMPLATE_DIR, render_docx
    from app.render_pdf import build_pdf

    _, mau = build_context({}, {}, {})
    ctx = {k: (v or f"[{k}]") for k, v in mau.items()}
    du = True

    for loai, thong_tin in spec.CONTRACTS.items():
        duong_dan = TEMPLATE_DIR / thong_tin["template"]

        # Template dùng biến mà code không bao giờ cấp -> chỗ đó in ra rỗng.
        dung = DocxTemplate(duong_dan).get_undeclared_template_variables()
        if (la := dung - set(ctx)):
            print(f"{LOI} {thong_tin['template']} dùng biến code không có: {sorted(la)}")
            du = False

        day_du = dien_cho_trong(ctx, loai)
        docx = render_docx(loai, day_du)
        buf = io.BytesIO()
        build_pdf([(loai, day_du)], buf)
        pdf = buf.getvalue()
        so_trang = len(re.findall(rb"/Type\s*/Page[^s]", pdf))

        if not pdf.startswith(b"%PDF") or so_trang < 1:
            print(f"{LOI} {loai}: PDF dựng ra không hợp lệ")
            du = False
        else:
            print(f"{OK} {loai}: Word {len(docx) // 1024}KB, PDF {len(pdf) // 1024}KB, {so_trang} trang")

    return du


# ───────────────────────── 5. cú pháp JavaScript ─────────────────────────

def kiem_tra_javascript():
    print("5. Cú pháp JavaScript")
    files = ["CONTRACTS/app/static/app.js", "CONTRACTS/app/static/backend-server.js",
             "web/backend-pyodide.js"]
    try:
        for f in files:
            r = subprocess.run(["node", "--check", str(GOC / f)], capture_output=True, text=True)
            if r.returncode:
                print(f"{LOI} {f}: {r.stderr.strip().splitlines()[-1]}")
                return False
    except FileNotFoundError:
        print("  – không có node, bỏ qua")
        return True
    print(f"{OK} {len(files)} file hợp lệ")
    return True


def main():
    kiem_tra = [kiem_tra_du_lieu_ca_nhan, kiem_tra_ban_web, kiem_tra_giao_dien,
                kiem_tra_dung_hop_dong, kiem_tra_javascript]
    ket_qua = []
    for ham in kiem_tra:
        ket_qua.append(ham())
        print()

    if all(ket_qua):
        print(f"Tất cả {len(ket_qua)} mục đều đạt.")
        return 0
    print(f"Có {ket_qua.count(False)}/{len(ket_qua)} mục không đạt.")
    return 1


if __name__ == "__main__":
    sys.exit(main())
