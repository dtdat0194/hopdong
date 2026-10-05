#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Dựng bản Word của hướng dẫn sử dụng từ chính trang huong-dan.html.

Chỉ có một bản gốc là trang web; chạy lại file này mỗi khi sửa trang để bản Word
khớp theo, khỏi phải nhớ sửa hai nơi.

    python3 CONTRACTS/scripts/xuat_huong_dan_word.py

Cần pandoc: brew install pandoc
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

GOC = Path(__file__).resolve().parents[2]
NGUON = GOC / "CONTRACTS/app/static/huong-dan.html"
# File Word chỉ dùng để gửi qua email, không đẩy lên kho mã (xem .gitignore).
DICH = GOC / "Huong_dan_su_dung.docx"

FONT, CO_CHU = "Times New Roman", 13

# Khung màu trên web không sang Word được, nên đổi thành trích dẫn kèm nhãn chữ.
NHAN_KHUNG = {"warn": "Lưu ý: ", "info": "Ghi chú: ", "ok": ""}


def chuan_bi_html(html):
    """Gọt trang web thành phần nội dung hợp với văn bản in."""
    than = re.search(r'<main class="doc">(.*?)</main>', html, re.S).group(1)

    # Mục lục bấm được chỉ có nghĩa trên web; Word tự có khung điều hướng riêng.
    than = re.sub(r'<div class="muc-luc">.*?</div>\s*</div>', "", than, flags=re.S)
    than = re.sub(r'<div class="muc-luc">.*?</div>', "", than, flags=re.S)

    # Danh sách các bước đang tự vẽ số bằng CSS; trả về thẻ ol thường để Word đánh số.
    than = than.replace('<ol class="buoc">', "<ol>")
    # Tiêu đề mỗi bước phải xuống dòng, không dính liền vào phần giải thích.
    than = re.sub(r"(<li>\s*)<strong>(.*?)</strong>", r"\1<strong>\2</strong><br>", than, flags=re.S)

    def doi_khung(m):
        nhan = NHAN_KHUNG.get(m.group(1), "")
        return f"<blockquote><p>{nhan}{m.group(2)}</p></blockquote>"

    than = re.sub(r'<div class="note (\w+)">(.*?)</div>', doi_khung, than, flags=re.S)
    than = re.sub(r'<p class="lead">(.*?)</p>', r"<p><em>\1</em></p>", than, flags=re.S)

    return ('<html><head><meta charset="utf-8"></head><body>'
            f"{than}</body></html>")


def ke_bang(bang, rong_cot):
    """Pandoc dựng bảng không viền và chia cột đều nhau; kẻ lại cho dễ đọc."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    vien = OxmlElement("w:tblBorders")
    for canh in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + canh)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "A6A6A6")
        vien.append(e)
    bang._tbl.tblPr.append(vien)

    bang.autofit = False
    for hang in bang.rows:
        for o, rong in zip(hang.cells, rong_cot):
            o.width = Cm(rong)

    for o in bang.rows[0].cells:         # hàng tiêu đề in đậm
        for p in o.paragraphs:
            for r in p.runs:
                r.bold = True


def dinh_dang(duong_dan):
    """Đặt lại phông chữ và căn lề cho hợp với văn bản nội bộ."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document(duong_dan)

    for ten in ("Normal", "Body Text", "Compact", "First Paragraph", "Block Text"):
        if ten in (s.name for s in doc.styles):
            st = doc.styles[ten]
            st.font.name, st.font.size = FONT, Pt(CO_CHU)
            if st.element.rPr is not None and st.element.rPr.rFonts is not None:
                st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)

    # Cột đầu chỉ chứa nhãn ngắn, phần giải thích dài nên để cột sau rộng hơn hẳn.
    for bang in doc.tables:
        ke_bang(bang, [4.3, 12.2])

    # Tiêu đề chung, chèn lên đầu trang.
    tieu_de = doc.paragraphs[0].insert_paragraph_before()
    tieu_de.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tieu_de.add_run("HƯỚNG DẪN SỬ DỤNG")
    r.bold, r.font.name, r.font.size = True, FONT, Pt(16)

    phu = doc.paragraphs[1].insert_paragraph_before()
    phu.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = phu.add_run("Công cụ xuất hợp đồng")
    r.italic, r.font.name, r.font.size = True, FONT, Pt(CO_CHU)
    phu.paragraph_format.space_after = Pt(18)

    doc.save(duong_dan)


def main():
    if not NGUON.exists():
        print(f"Không thấy {NGUON}")
        return 1
    try:
        subprocess.run(["pandoc", "--version"], capture_output=True, check=True)
    except (FileNotFoundError, subprocess.CalledProcessError):
        print("Cần pandoc: brew install pandoc")
        return 1

    html = chuan_bi_html(NGUON.read_text(encoding="utf-8"))
    with tempfile.NamedTemporaryFile("w", suffix=".html", encoding="utf-8", delete=False) as f:
        f.write(html)
        tam = f.name

    subprocess.run(["pandoc", tam, "-f", "html", "-t", "docx", "-o", str(DICH)], check=True)
    Path(tam).unlink()
    dinh_dang(DICH)

    print(f"Đã dựng {DICH.relative_to(GOC)} ({DICH.stat().st_size // 1024}KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
